"""
Build Phase II Final Report DOCX with professional Cover Page.
No manual merging needed - one clean file.
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

BASE = Path(r"C:\Users\Kelvin Lam\Documents\GitHub\FYP\FYP\Documents\canvasdoucment")
MD_FILE = BASE / "Phase_II_Final_Report.md"
OUT_FILE = BASE / "Phase_II_Final_Report_v2.docx"

FONT = "Times New Roman"

# ─── Cover Page Data ───
COVER = {
    "university": "City University of Hong Kong",
    "school": "School of Creative Media",
    "programme": "BScCM — Bachelor of Science in Creative Media",
    "course": "SM4712B (LA2) Graduation Thesis/Project",
    "subtitle": "BScCM FYP Implementation — Phase II",
    "title": "Arcane Souls: Rebirth",
    "group": "AIGEN",
    "advisor_scm": "Jussi Pekka HOLOPAINEN",
    "advisor_cs": "—",
    "students": [
        ("Lee Chun Kit", "57306141", "cklee96", "Group Leader"),
        ("Lam Chi Him", "57185861", "kelvelam6", "Member"),
    ],
    "date": "12 April 2026",
}


def set_run(run, text, size=12, bold=False, color=None, font_name=FONT):
    """Helper to set run properties."""
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    # Force East-Asian font too
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = r.makeelement(qn('w:rPr'), {})
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)


def add_cover_line(doc, text, size=12, bold=False, color=None, spacing_after=6, alignment=WD_ALIGN_PARAGRAPH.CENTER):
    """Add a single centered line to doc."""
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(spacing_after)
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)
    return p


def build_cover(doc):
    """Build a professional cover page."""
    # Top spacing
    for _ in range(3):
        add_cover_line(doc, "", size=12, spacing_after=0)

    # University name
    add_cover_line(doc, COVER["university"], size=22, bold=True, spacing_after=4)
    
    # School
    add_cover_line(doc, COVER["school"], size=16, bold=True, spacing_after=2)
    
    # Programme
    add_cover_line(doc, COVER["programme"], size=13, spacing_after=20)
    
    # Divider line
    add_cover_line(doc, "─" * 50, size=10, color=(150, 150, 150), spacing_after=20)
    
    # Course
    add_cover_line(doc, COVER["course"], size=13, bold=True, spacing_after=4)
    add_cover_line(doc, COVER["subtitle"], size=12, spacing_after=30)
    
    # Main title (big)
    add_cover_line(doc, COVER["title"], size=28, bold=True, spacing_after=30)
    
    # Divider
    add_cover_line(doc, "─" * 50, size=10, color=(150, 150, 150), spacing_after=24)
    
    # Group
    add_cover_line(doc, f"Group: {COVER['group']}", size=14, bold=True, spacing_after=16)
    
    # Advisor
    add_cover_line(doc, "Advisor (SCM)", size=11, bold=True, spacing_after=2)
    add_cover_line(doc, COVER["advisor_scm"], size=13, spacing_after=16)
    
    # Students header
    add_cover_line(doc, "Students", size=11, bold=True, spacing_after=6)
    
    # Student rows
    for name, sid, eid, role in COVER["students"]:
        add_cover_line(doc, f"{name}  (SID: {sid},  EID: {eid})  —  {role}", size=12, spacing_after=4)
    
    # Spacing before date
    add_cover_line(doc, "", size=12, spacing_after=20)
    
    # Date
    add_cover_line(doc, f"Submission Date: {COVER['date']}", size=13, bold=True, spacing_after=4)

    # Blog links
    add_cover_line(doc, "", size=8, spacing_after=10)
    add_cover_line(doc, "Individual Blogs:", size=10, bold=True, spacing_after=2)
    add_cover_line(doc, "Lee Chun Kit: https://leechunkit01255210.wixsite.com/e-portfolio", size=9, spacing_after=2)
    add_cover_line(doc, "Lam Chi Him: https://nihonjin864.wixsite.com/lamchihim", size=9, spacing_after=0)
    
    # Page break after cover
    doc.add_page_break()


# ─── Markdown → DOCX body ───

def add_styled_paragraph(doc, text, style_name, size=12, bold=False):
    p = doc.add_paragraph(style=style_name)
    for run in p.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
    if not p.runs:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
    return p


def build_body(doc, md_text):
    """Parse markdown and add content to doc."""
    lines = md_text.split("\n")
    in_code = False
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            return
        # Filter separator rows
        data = [r for r in table_rows if not all(c.strip().replace("-", "") == "" for c in r)]
        if len(data) < 1:
            table_rows = []
            in_table = False
            return
        cols = len(data[0])
        tbl = doc.add_table(rows=len(data), cols=cols, style="Table Grid")
        for i, row_data in enumerate(data):
            for j in range(cols):
                cell_text = row_data[j].strip() if j < len(row_data) else ""
                cell = tbl.cell(i, j)
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(cell_text)
                run.font.name = FONT
                run.font.size = Pt(9)
                run.font.bold = (i == 0)
        table_rows = []
        in_table = False

    for line in lines:
        stripped = line.strip()

        # Skip the very first metadata lines we already put in cover
        if stripped.startswith("# BScCM") or stripped.startswith("# Phase II Final"):
            continue
        if stripped.startswith("**City University") or stripped.startswith("**School of") or stripped.startswith("**Department"):
            continue
        if stripped.startswith("- **Project Title") or stripped.startswith("- **Course Code") or stripped.startswith("- **Group Name"):
            continue
        if stripped.startswith("- **Advisor") or stripped.startswith("- **Students") or stripped.startswith("- **Submission"):
            continue
        if stripped.startswith("- **Individual"):
            continue
        if stripped.startswith("  - Lee Chun Kit:") or stripped.startswith("  - Lam Chi Him:"):
            continue
        if stripped.startswith("  - Lee Chun Kit (SID") or stripped.startswith("  - Lam Chi Him (SID"):
            continue

        # Code blocks
        if stripped.startswith("```"):
            if in_table:
                flush_table()
            in_code = not in_code
            if in_code:
                continue
            else:
                continue

        if in_code:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            continue

        # Tables
        if "|" in stripped and stripped.startswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            table_rows.append(cells)
            continue
        else:
            if in_table:
                flush_table()

        # Headings
        if stripped.startswith("####"):
            add_styled_paragraph(doc, stripped.lstrip("#").strip(), "Heading 4", size=11, bold=True)
            continue
        if stripped.startswith("###"):
            add_styled_paragraph(doc, stripped.lstrip("#").strip(), "Heading 3", size=12, bold=True)
            continue
        if stripped.startswith("##"):
            add_styled_paragraph(doc, stripped.lstrip("#").strip(), "Heading 2", size=14, bold=True)
            continue
        if stripped.startswith("#"):
            add_styled_paragraph(doc, stripped.lstrip("#").strip(), "Heading 1", size=16, bold=True)
            continue

        # Horizontal rules
        if stripped == "---":
            continue

        # Empty lines
        if stripped == "":
            continue

        # Bullet points
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            text = stripped[2:]
            # Handle bold within bullet
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = FONT
                run.font.size = Pt(11)
            continue

        # Numbered lists
        m = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            text = m.group(2)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = FONT
                run.font.size = Pt(11)
            continue

        # Normal paragraphs
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.font.bold = True
            else:
                run = p.add_run(part)
            run.font.name = FONT
            run.font.size = Pt(11)

    # Flush any remaining table
    if in_table:
        flush_table()


def main():
    md_text = MD_FILE.read_text(encoding="utf-8")
    print(f"Reading MD: {MD_FILE}")

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)

    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    print("Building cover page...")
    build_cover(doc)

    print("Building body content...")
    build_body(doc, md_text)

    doc.save(str(OUT_FILE))
    print(f"\n=== DONE ===")
    print(f"Saved: {OUT_FILE}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
