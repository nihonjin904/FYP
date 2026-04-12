#!/usr/bin/env python3
"""Generate Phase II Final Report DOCX from markdown."""

from docx import Document
from docx.shared import Pt
import re

def clean_md(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    return text.strip()

def add_bold_text(paragraph, text):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)

def main():
    doc = Document()

    # Default style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # Read markdown
    import os
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(script_dir, 'canvasdoucment', 'Phase_II_Final_Report.md')
    if not os.path.exists(md_path):
        md_path = os.path.join(script_dir, 'Phase_II_Final_Report.md')
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    i = 0
    in_code = False
    table_rows = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code block toggle
        if stripped.startswith('```'):
            in_code = not in_code
            i += 1
            continue

        # Code content
        if in_code:
            p = doc.add_paragraph()
            run = p.add_run(line.rstrip())
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            i += 1
            continue

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == '---':
            p = doc.add_paragraph()
            p.add_run('_' * 60).font.size = Pt(8)
            i += 1
            continue

        # Table handling
        if '|' in stripped and stripped.startswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            # Skip separator rows
            if all(set(c.replace(' ', '')) <= set('-:') for c in cells):
                i += 1
                continue
            table_rows.append(cells)
            # Peek next line
            next_is_table = (i + 1 < len(lines) and
                             '|' in lines[i + 1].strip() and
                             lines[i + 1].strip().startswith('|'))
            if not next_is_table:
                # Flush table
                if table_rows:
                    max_cols = max(len(r) for r in table_rows)
                    tbl = doc.add_table(rows=len(table_rows), cols=max_cols)
                    tbl.style = 'Table Grid'
                    for ri, row in enumerate(table_rows):
                        for ci, cell in enumerate(row):
                            if ci < max_cols:
                                cell_text = clean_md(cell)
                                tbl.rows[ri].cells[ci].text = cell_text
                                for para in tbl.rows[ri].cells[ci].paragraphs:
                                    for run in para.runs:
                                        run.font.size = Pt(10)
                                        run.font.name = 'Times New Roman'
                    # Bold header row
                    for ci in range(max_cols):
                        for para in tbl.rows[0].cells[ci].paragraphs:
                            for run in para.runs:
                                run.bold = True
                    doc.add_paragraph()
                table_rows = []
            i += 1
            continue

        # Headings
        if stripped.startswith('#### '):
            doc.add_heading(clean_md(stripped[5:]), level=4)
            i += 1
            continue
        elif stripped.startswith('### '):
            doc.add_heading(clean_md(stripped[4:]), level=3)
            i += 1
            continue
        elif stripped.startswith('## '):
            doc.add_heading(clean_md(stripped[3:]), level=2)
            i += 1
            continue
        elif stripped.startswith('# ') and not stripped.startswith('## '):
            doc.add_heading(clean_md(stripped[2:]), level=1)
            i += 1
            continue

        # Bullet points
        if stripped.startswith('- '):
            text = clean_md(stripped[2:])
            doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue

        # Numbered items
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            text = clean_md(num_match.group(2))
            doc.add_paragraph(text, style='List Number')
            i += 1
            continue

        # Regular text
        text = clean_md(stripped)
        if text:
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            add_bold_text(p, stripped)

        i += 1

    output = os.path.join(script_dir, 'canvasdoucment', 'Phase_II_Final_Report.docx')
    doc.save(output)
    print(f'Saved: {output}')
    print(f'Paragraphs: {len(doc.paragraphs)}')
    print(f'Tables: {len(doc.tables)}')

if __name__ == '__main__':
    main()
